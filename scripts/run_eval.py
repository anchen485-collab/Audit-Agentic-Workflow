import argparse
import json
import sys
from pathlib import Path
from typing import Any

from app.domain.models import AuditCase
from app.infrastructure.evidence_provider import MockCompanyEvidenceProvider
from app.infrastructure.excel_reader import read_audit_excel
from app.infrastructure.rule_repository import load_category_rules
from app.infrastructure.settings import get_settings
from app.services.case_service import create_audit_cases
from app.services.decision_service import calculate_decision_accuracy
from app.workflow.audit_workflow import build_audit_workflow

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CASES_FILE = PROJECT_ROOT / "evals" / "golden_cases.json"
FALLBACK_CASES_FILE = PROJECT_ROOT / "tests" / "fixtures" / "sample_cases.json"
DEFAULT_RULES_FILE = PROJECT_ROOT / "data" / "rules" / "category_rules.json"
DEFAULT_EVIDENCE_FILE = PROJECT_ROOT / "data" / "evidence" / "mock_company_evidence.json"
EXPECTED_ACTION_HEADERS = {"expected_action", "期望动作", "标准答案", "人工标注动作"}


class EvaluationFailedError(AssertionError):
    """评测未达到业务门槛时抛出，携带错例方便定位。"""

    def __init__(self, accuracy: float, min_accuracy: float, mismatches: list[dict[str, Any]]):
        super().__init__(f"accuracy {accuracy:.4f} is below min_accuracy {min_accuracy:.4f}")
        self.accuracy = accuracy
        self.min_accuracy = min_accuracy
        self.mismatches = mismatches


class EvaluationDataError(ValueError):
    """评测数据不适合当前评测模式时抛出，避免生成误导性指标。"""


def run_eval(
    cases_file: Path = DEFAULT_CASES_FILE,
    excel_file: Path | None = None,
    rules_file: Path = DEFAULT_RULES_FILE,
    evidence_file: Path = DEFAULT_EVIDENCE_FILE,
    evidence_provider: Any | None = None,
    classification_provider: Any | None = None,
    use_real_providers: bool = True,
    min_accuracy: float = 0.8,
) -> dict[str, Any]:
    """运行审批 Agent 评测，并返回结构化指标。"""
    if not cases_file.exists():
        cases_file = FALLBACK_CASES_FILE

    audit_cases, expected_actions = _load_eval_cases(cases_file=cases_file, excel_file=excel_file)
    if use_real_providers:
        _validate_real_eval_cases(audit_cases)

    current_evidence_provider = evidence_provider or _build_evidence_provider(
        evidence_file=evidence_file,
        use_real_providers=use_real_providers,
    )
    current_classification_provider = classification_provider or _build_classification_provider(
        use_real_providers=use_real_providers,
    )

    workflow = build_audit_workflow()
    result = workflow.invoke(
        {
            "cases": audit_cases,
            "rules": load_category_rules(rules_file),
            "evidence_provider": current_evidence_provider,
            "classification_provider": current_classification_provider,
            "expected_actions": expected_actions,
        }
    )

    decision_models = result["decisions"]
    decisions = [decision.model_dump(mode="json") for decision in decision_models]
    evidence_packages = result.get("evidence_packages", {})
    accuracy_context = _build_accuracy_context(
        decisions=decision_models,
        expected_actions=expected_actions,
        evidence_packages=evidence_packages,
    )
    mismatches = _build_mismatches(decisions, accuracy_context["eligible_expected_actions"])
    metrics = dict(result["metrics"])
    metrics.update(accuracy_context["metrics"])
    eval_result = {
        "case_count": len(audit_cases),
        "decision_count": len(result["decisions"]),
        "review_task_count": len(result["review_tasks"]),
        "metrics": metrics,
        "mismatches": mismatches,
        "decisions": decisions,
        "review_tasks": [task.model_dump(mode="json") for task in result["review_tasks"]],
    }

    accuracy = float(metrics.get("accuracy", 0.0))
    has_expected_actions = bool(accuracy_context["eligible_expected_actions"])
    if has_expected_actions and accuracy < min_accuracy:
        raise EvaluationFailedError(accuracy=accuracy, min_accuracy=min_accuracy, mismatches=mismatches)

    return eval_result


def _build_mismatches(
    decisions: list[dict[str, Any]],
    expected_actions: dict[str, str],
) -> list[dict[str, Any]]:
    """生成错例列表，方便定位业务评测失败原因。"""
    mismatches = []
    for decision in decisions:
        case_id = decision["case_id"]
        if case_id not in expected_actions:
            continue

        expected_action = expected_actions.get(case_id)
        actual_action = decision["action"]
        if expected_action == actual_action:
            continue

        mismatches.append(
            {
                "case_id": case_id,
                "expected_action": expected_action,
                "actual_action": actual_action,
                "reason": decision.get("reason", ""),
                "suggestion": decision.get("suggestion", ""),
            }
        )

    return mismatches


def _load_eval_cases(
    cases_file: Path,
    excel_file: Path | None,
) -> tuple[list[AuditCase], dict[str, str] | None]:
    """读取 JSON 或 Excel 评测集，并返回审批单与可选标准答案。"""
    if excel_file is not None:
        rows = read_audit_excel(excel_file)
        audit_cases = create_audit_cases(rows, batch_id="eval_excel")
        expected_actions = _load_expected_actions_from_excel(excel_file, audit_cases)
        return audit_cases, expected_actions or None

    raw_cases = json.loads(cases_file.read_text(encoding="utf-8"))
    audit_cases = _build_audit_cases(raw_cases)
    expected_actions = {
        item["case_id"]: item["expected_action"]
        for item in raw_cases
        if item.get("expected_action")
    }
    return audit_cases, expected_actions or None


def _load_expected_actions_from_excel(excel_file: Path, audit_cases: list[AuditCase]) -> dict[str, str]:
    """从 Excel 的标准答案列读取 expected_action；没有该列时返回空字典。"""
    from openpyxl import load_workbook

    workbook = load_workbook(excel_file, data_only=True)
    sheet = workbook.active
    header_cells = next(sheet.iter_rows(min_row=1, max_row=1))
    header_map = {
        _cell_to_text(cell.value): index
        for index, cell in enumerate(header_cells)
        if _cell_to_text(cell.value)
    }
    expected_index = next(
        (index for header, index in header_map.items() if header in EXPECTED_ACTION_HEADERS),
        None,
    )
    if expected_index is None:
        return {}

    case_by_row_number = {
        audit_case.row_number: audit_case
        for audit_case in audit_cases
    }
    expected_actions = {}
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        audit_case = case_by_row_number.get(row_number)
        if audit_case is None:
            continue

        action = _cell_to_text(cells[expected_index].value if expected_index < len(cells) else "")
        if action:
            expected_actions[audit_case.case_id] = action

    return expected_actions


def _build_audit_cases(raw_cases: list[dict[str, Any]]) -> list[AuditCase]:
    """把 golden cases 转成 workflow 使用的 AuditCase。"""
    audit_cases = []

    for row_number, item in enumerate(raw_cases, start=2):
        input_data = item["input"]
        audit_cases.append(
            AuditCase(
                case_id=item["case_id"],
                batch_id="eval_batch",
                row_number=row_number,
                employee_name=input_data["employee_name"],
                company_raw=input_data["company_raw"],
                company_name=input_data["company_name"],
                website_url=input_data.get("website_url", ""),
                entered_level1=input_data["entered_level1"],
                entered_subcategory=input_data["entered_subcategory"],
                entered_stage=input_data.get("entered_stage", ""),
                status="pending",
            )
        )

    return audit_cases


def _validate_real_eval_cases(audit_cases: list[AuditCase]) -> None:
    """真实流程评测必须使用真实官网 URL，不能用 example.com 这类 mock 域名。"""
    invalid_case_ids = []
    for audit_case in audit_cases:
        if _is_mock_website_url(audit_case.website_url):
            invalid_case_ids.append(audit_case.case_id)

    if invalid_case_ids:
        preview = ", ".join(invalid_case_ids[:5])
        raise EvaluationDataError(
            "真实流程 eval 发现 mock 官网 URL，无法生成有业务意义的准确率。"
            f"请把这些 case 换成真实可访问官网，或使用 --mock 做离线回归。case_id={preview}"
        )


def _is_mock_website_url(website_url: str) -> bool:
    """识别文档和测试中常用的占位官网域名。"""
    normalized_url = website_url.lower()
    return any(
        marker in normalized_url
        for marker in {
            "example.com",
            "example.org",
            "example.net",
            ".example.",
        }
    )


def _build_accuracy_context(
    decisions: list[Any],
    expected_actions: dict[str, str] | None,
    evidence_packages: dict[str, Any],
) -> dict[str, Any]:
    """过滤官网不可访问样本后计算准确率上下文。"""
    if expected_actions is None:
        return {
            "eligible_expected_actions": {},
            "metrics": {
                "accuracy_eligible_count": 0,
                "accuracy_excluded_count": 0,
            },
        }

    excluded_case_ids = {
        case_id
        for case_id, evidence in evidence_packages.items()
        if _is_external_evidence_unavailable(evidence)
    }
    eligible_expected_actions = {
        case_id: action
        for case_id, action in expected_actions.items()
        if case_id not in excluded_case_ids
    }
    metrics = {
        "accuracy_eligible_count": len(eligible_expected_actions),
        "accuracy_excluded_count": len(set(expected_actions) & excluded_case_ids),
    }
    if eligible_expected_actions:
        eligible_decisions = [
            decision
            for decision in decisions
            if decision.case_id in eligible_expected_actions
        ]
        metrics["accuracy"] = calculate_decision_accuracy(
            decisions=eligible_decisions,
            expected_actions=eligible_expected_actions,
        )

    return {
        "eligible_expected_actions": eligible_expected_actions,
        "metrics": metrics,
    }


def _is_external_evidence_unavailable(evidence: Any) -> bool:
    """判断证据不可用是否属于外部官网访问问题。"""
    errors = getattr(evidence, "errors", []) or []
    return any(
        "官网访问失败" in error
        or "缺少企业官网 URL" in error
        or "缺少官网 URL" in error
        for error in errors
    )


def _cell_to_text(value: Any) -> str:
    """把 Excel 单元格值转成字符串。"""
    if value is None:
        return ""
    return str(value).strip()


def _build_evidence_provider(evidence_file: Path, use_real_providers: bool):
    """构造证据 provider；真实评测默认使用 .env 中配置的官网爬虫。"""
    if not use_real_providers:
        return MockCompanyEvidenceProvider(evidence_file)

    from app.api.dependencies import get_evidence_provider

    return get_evidence_provider()


def _build_classification_provider(use_real_providers: bool):
    """构造分类 provider；真实评测默认使用 .env 中配置的 LLM。"""
    if not use_real_providers:
        return None

    from app.api.dependencies import get_classification_provider

    settings = get_settings()
    if not settings.llm_enabled:
        raise ValueError("真实评测必须配置 LLM_ENABLED=true，否则无法验证 LLM 语义判断能力。")

    return get_classification_provider()


def main() -> None:
    """命令行入口。"""
    parser = argparse.ArgumentParser(description="运行审批 Agent 评测")
    parser.add_argument("--cases-file", type=Path, default=DEFAULT_CASES_FILE)
    parser.add_argument("--excel-file", type=Path, default=None)
    parser.add_argument("--rules-file", type=Path, default=DEFAULT_RULES_FILE)
    parser.add_argument("--evidence-file", type=Path, default=DEFAULT_EVIDENCE_FILE)
    parser.add_argument("--min-accuracy", type=float, default=0.8)
    parser.add_argument(
        "--mock",
        action="store_true",
        help="使用 mock 证据和 mock 分类 provider 做稳定回归；默认读取 .env 走真实 provider。",
    )
    parser.add_argument(
        "--use-llm",
        action="store_true",
        help="兼容旧参数；当前默认已经使用真实 LLM provider。",
    )
    args = parser.parse_args()

    try:
        result = run_eval(
            cases_file=args.cases_file,
            excel_file=args.excel_file,
            rules_file=args.rules_file,
            evidence_file=args.evidence_file,
            use_real_providers=not args.mock,
            min_accuracy=args.min_accuracy,
        )
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except EvaluationDataError as error:
        print(
            json.dumps(
                {
                    "error": str(error),
                    "error_type": "evaluation_data_error",
                },
                ensure_ascii=False,
                indent=2,
            )
        )
        sys.exit(1)
    except EvaluationFailedError as error:
        print(
            json.dumps(
                {
                    "error": str(error),
                    "accuracy": error.accuracy,
                    "min_accuracy": error.min_accuracy,
                    "mismatches": error.mismatches,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
        sys.exit(1)


if __name__ == "__main__":
    main()
