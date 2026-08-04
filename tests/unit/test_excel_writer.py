from openpyxl import load_workbook

from app.domain.decisions import DecisionAction, DecisionStatus
from app.domain.models import AgentAssessment, ApprovalDecision, AuditCase, EvidencePackage, EvidenceSnippet, ReviewTask
from app.infrastructure.excel_writer import write_audit_result_excel


def test_write_audit_result_excel_creates_three_sheets(tmp_path):
    """导出结果 Excel 应该包含三个 Sheet。"""
    output_path = tmp_path / "result.xlsx"
    cases = [
        AuditCase(
            case_id="case_001",
            batch_id="batch_001",
            row_number=2,
            employee_name="张三",
            company_raw="杭州云启软件有限公司",
            company_name="杭州云启软件有限公司",
            website_url="https://example.com",
            entered_level1="信软技术服务",
            entered_level2="软件和信息技术",
            entered_subcategory="软件开发",
            status="completed",
        )
    ]
    decisions = [
        ApprovalDecision(
            case_id="case_001",
            action=DecisionAction.APPROVED,
            status=DecisionStatus.COMPLETED,
            confidence=90,
            risk_score=10,
            reason="证据支持员工录入分类。",
            suggestion="",
            needs_review=False,
        )
    ]
    review_tasks = [
        ReviewTask(
            task_id="review_001",
            case_id="case_002",
            priority="high",
            reason="分类不存在。",
            status="pending",
        )
    ]

    write_audit_result_excel(output_path, cases, decisions, review_tasks)

    workbook = load_workbook(output_path)

    assert workbook.sheetnames == ["审批明细", "人员汇总", "人工复核任务"]


def test_write_audit_result_excel_writes_detail_rows(tmp_path):
    """审批明细应该写入 case 和 decision 信息。"""
    output_path = tmp_path / "result.xlsx"
    cases = [
        AuditCase(
            case_id="case_001",
            batch_id="batch_001",
            row_number=2,
            employee_name="张三",
            company_raw="杭州云启软件有限公司",
            company_name="杭州云启软件有限公司",
            website_url="https://example.com",
            entered_level1="信软技术服务",
            entered_level2="软件和信息技术",
            entered_subcategory="软件开发",
            status="completed",
        )
    ]
    decisions = [
        ApprovalDecision(
            case_id="case_001",
            action=DecisionAction.APPROVED,
            status=DecisionStatus.COMPLETED,
            confidence=90,
            risk_score=10,
            reason="证据支持员工录入分类。",
            suggestion="",
            needs_review=False,
        )
    ]

    write_audit_result_excel(output_path, cases, decisions, [])

    workbook = load_workbook(output_path)
    sheet = workbook["审批明细"]

    headers = [cell.value for cell in sheet[1]]

    assert sheet["A1"].value == "审批单ID"
    assert sheet["A2"].value == "case_001"
    assert "录入人员" in headers
    assert "员工录入一级分类" in headers
    assert "员工录入二级分类" in headers
    assert "员工录入细分" in headers
    assert "企业名称" in headers
    assert "企业官网" in headers
    assert sheet.cell(row=2, column=headers.index("录入人员") + 1).value == "张三"
    assert sheet.cell(row=2, column=headers.index("员工录入一级分类") + 1).value == "信软技术服务"
    assert sheet.cell(row=2, column=headers.index("员工录入二级分类") + 1).value == "软件和信息技术"
    assert sheet.cell(row=2, column=headers.index("员工录入细分") + 1).value == "软件开发"
    assert sheet.cell(row=2, column=headers.index("企业名称") + 1).value == "杭州云启软件有限公司"
    assert sheet.cell(row=2, column=headers.index("企业官网") + 1).value == "https://example.com"
    assert sheet.cell(row=2, column=headers.index("审批动作") + 1).value == "approved"
    assert sheet.cell(row=2, column=headers.index("置信度") + 1).value == 90


def test_write_audit_result_excel_writes_evidence_columns(tmp_path):
    """审批明细应该写入证据来源、质量分和证据片段。"""
    output_path = tmp_path / "result.xlsx"
    cases = [
        AuditCase(
            case_id="case_001",
            batch_id="batch_001",
            row_number=2,
            employee_name="张三",
            company_raw="杭州云启软件有限公司",
            company_name="杭州云启软件有限公司",
            website_url="https://example.com",
            entered_level1="信软技术服务",
            entered_level2="软件和信息技术",
            entered_subcategory="软件开发",
            status="completed",
        )
    ]
    decisions = [
        ApprovalDecision(
            case_id="case_001",
            action=DecisionAction.APPROVED,
            status=DecisionStatus.COMPLETED,
            confidence=90,
            risk_score=10,
            reason="证据支持员工录入分类。",
            suggestion="",
            needs_review=False,
        )
    ]
    evidence_packages = {
        "case_001": EvidencePackage(
            case_id="case_001",
            company_name="杭州云启软件有限公司",
            sources=["website"],
            business_scope="公司提供软件开发服务。",
            snippets=[EvidenceSnippet(source="website", text="公司提供软件开发服务。", url="https://example.com")],
            quality_score=90,
            errors=[],
        )
    }

    write_audit_result_excel(output_path, cases, decisions, [], evidence_packages=evidence_packages)

    workbook = load_workbook(output_path)
    sheet = workbook["审批明细"]
    headers = [cell.value for cell in sheet[1]]

    assert "证据来源" in headers
    assert "证据质量分" in headers
    assert "证据片段" in headers
    assert sheet.cell(row=2, column=headers.index("证据来源") + 1).value == "website"
    assert sheet.cell(row=2, column=headers.index("证据质量分") + 1).value == 90


def test_write_audit_result_excel_writes_classification_agent_error_columns(tmp_path):
    """结果 Excel 应该展示 Classification Agent 的真实错误摘要。"""
    output_path = tmp_path / "result.xlsx"
    cases = [
        AuditCase(
            case_id="case_001",
            batch_id="batch_001",
            row_number=2,
            employee_name="张三",
            company_raw="测试公司",
            company_name="测试公司",
            website_url="https://example.com",
            entered_level1="信软技术服务",
            entered_level2="软件和信息技术",
            entered_subcategory="软件开发",
            status="completed",
        )
    ]
    decisions = [
        ApprovalDecision(
            case_id="case_001",
            action=DecisionAction.NEEDS_MORE_EVIDENCE,
            status=DecisionStatus.PENDING_REVIEW,
            confidence=80,
            risk_score=70,
            reason="外部证据不足：官网访问失败：502 Bad Gateway",
            suggestion="请补充证据。",
            needs_review=True,
        )
    ]
    assessments = {
        "case_001": [
            AgentAssessment(
                agent_name="classification_agent",
                status="success",
                confidence=80,
                findings=[],
                output={
                    "case_id": "case_001",
                    "evidence_error_type": "evidence_unavailable",
                    "normalization_warning": "",
                },
                duration_ms=1,
            )
        ]
    }

    write_audit_result_excel(output_path, cases, decisions, [], assessments=assessments)

    sheet = load_workbook(output_path)["审批明细"]
    headers = [cell.value for cell in sheet[1]]

    assert "Classification Agent 状态" in headers
    assert "Classification Agent 错误原因" in headers
    assert "证据不足类型" in headers
    assert sheet.cell(row=2, column=headers.index("Classification Agent 状态") + 1).value == "success"
    assert sheet.cell(row=2, column=headers.index("证据不足类型") + 1).value == "evidence_unavailable"


def test_write_audit_result_excel_writes_rule_recall_granularity_columns(tmp_path):
    """结果 Excel 应该展示规则召回粒度和候选分类路径，方便定位分类表问题。"""
    output_path = tmp_path / "result.xlsx"
    cases = [
        AuditCase(
            case_id="case_001",
            batch_id="batch_001",
            row_number=2,
            employee_name="张三",
            company_raw="测试医院",
            company_name="测试医院",
            website_url="https://example.com",
            entered_level1="卫生和社会工作",
            entered_subcategory="医院",
            status="completed",
        )
    ]
    decisions = [
        ApprovalDecision(
            case_id="case_001",
            action=DecisionAction.APPROVED,
            status=DecisionStatus.COMPLETED,
            confidence=90,
            risk_score=10,
            reason="员工录入为二级分类，外部证据支持该二级分类。",
            suggestion="",
            needs_review=False,
        )
    ]
    assessments = {
        "case_001": [
            AgentAssessment(
                agent_name="rule_agent",
                status="success",
                confidence=90,
                findings=[],
                output={
                    "case_id": "case_001",
                    "entered_subcategory_granularity": "level2",
                    "granularity_reason": "员工录入细分命中二级分类。",
                    "category_candidate_paths": ["卫生和社会工作 / 医院"],
                },
                duration_ms=1,
            )
        ]
    }

    write_audit_result_excel(output_path, cases, decisions, [], assessments=assessments)

    sheet = load_workbook(output_path)["审批明细"]
    headers = [cell.value for cell in sheet[1]]

    assert "细分粒度" in headers
    assert "候选分类" in headers
    assert "召回原因" in headers
    assert sheet.cell(row=2, column=headers.index("细分粒度") + 1).value == "level2"
    assert sheet.cell(row=2, column=headers.index("候选分类") + 1).value == "卫生和社会工作 / 医院"


def test_write_audit_result_excel_uses_chinese_headers_for_all_sheets(tmp_path):
    """结果 Excel 的所有 Sheet 表头都应该使用中文。"""
    output_path = tmp_path / "result.xlsx"
    cases = [
        AuditCase(
            case_id="case_001",
            batch_id="batch_001",
            row_number=2,
            employee_name="张三",
            company_raw="杭州云启软件有限公司",
            company_name="杭州云启软件有限公司",
            website_url="https://example.com",
            entered_level1="信软技术服务",
            entered_level2="软件和信息技术",
            entered_subcategory="软件开发",
            status="completed",
        )
    ]
    decisions = [
        ApprovalDecision(
            case_id="case_001",
            action=DecisionAction.NEEDS_HUMAN_REVIEW,
            status=DecisionStatus.PENDING_REVIEW,
            confidence=70,
            risk_score=80,
            reason="需要人工确认。",
            suggestion="请人工确认。",
            needs_review=True,
        )
    ]
    review_tasks = [
        ReviewTask(
            task_id="review_001",
            case_id="case_001",
            priority="high",
            reason="需要人工确认。",
            status="pending",
        )
    ]

    write_audit_result_excel(output_path, cases, decisions, review_tasks)

    workbook = load_workbook(output_path)

    assert [cell.value for cell in workbook["审批明细"][1]][:6] == [
        "审批单ID",
        "批次ID",
        "原始行号",
        "录入人员",
        "企业名称",
        "企业官网",
    ]
    assert [cell.value for cell in workbook["人员汇总"][1]] == [
        "录入人员",
        "总数",
        "自动通过",
        "自动驳回",
        "证据不足",
        "人工复核",
    ]
    assert [cell.value for cell in workbook["人工复核任务"][1]] == [
        "复核任务ID",
        "审批单ID",
        "优先级",
        "复核原因",
        "任务状态",
        "审核员",
        "人工最终动作",
        "人工复核时间",
    ]
