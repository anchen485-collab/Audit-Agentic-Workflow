from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.app import create_app
from app.api.dependencies import get_app_settings
from app.infrastructure.batch_repository import list_file_access_audit_log_payloads
from app.infrastructure.settings import Settings


@pytest.fixture
def test_client(tmp_path):
    """创建使用临时 storage 的测试客户端。"""
    app = create_app()

    def override_settings() -> Settings:
        return Settings(AUDIT_STORAGE_DIR=tmp_path, LLM_ENABLED=False, _env_file=None)

    app.dependency_overrides[get_app_settings] = override_settings

    yield TestClient(app)

    app.dependency_overrides.clear()

def test_upload_audit_excel_returns_pipeline_result(tmp_path):
    """上传员工 Excel 后，应该返回审批结果和指标。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    _create_sample_excel(excel_path)

    client = TestClient(create_app())

    with excel_path.open("rb") as file:
        response = client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200


    data = response.json()
    assert data["case_count"] == 1
    assert data["decision_count"] == 1
    assert "metrics" in data
    assert data["metrics"]["decision_count"] == 1
    assert data["cases"][0]["employee_name"] == "张三"
    assert data["cases"][0]["entered_level1"] == "信软技术服务"
    assert data["cases"][0]["entered_subcategory"] == "软件开发"
    assert data["cases"][0]["company_name"] == "杭州云启软件有限公司"
    assert data["cases"][0]["website_url"] == "https://yunqi.example.com"
    assert data["cases"][0]["entered_subcategory_granularity"]
    assert "granularity_reason" in data["cases"][0]
    assert "category_candidate_paths" in data["cases"][0]
    assert data["decisions"][0]["case_id"].startswith(data["batch_id"])
    assert "result_file" not in data
    assert data["download_url"].startswith("/api/audit/batches/")


def test_audit_workbench_page_is_available():
    """应该能打开前端审批工作台页面。"""
    client = TestClient(create_app())

    response = client.get("/")

    assert response.status_code == 200
    assert "审批工作台" in response.text
    assert "/static/app.js" in response.text


def test_llm_health_returns_mock_status_when_llm_disabled(test_client):
    """LLM 关闭时，健康检查应该返回 mock provider 状态。"""
    response = test_client.get("/health/llm")

    assert response.status_code == 200
    data = response.json()
    assert data["enabled"] is False
    assert data["ok"] is True
    assert data["provider"] == "mock_classification"


def test_llm_health_does_not_expose_api_key_when_check_disabled(tmp_path):
    """LLM 健康检查不应该泄漏 API Key。"""
    app = create_app()

    def override_settings() -> Settings:
        return Settings(
            AUDIT_STORAGE_DIR=tmp_path,
            LLM_ENABLED=True,
            LLM_API_KEY="sk-secret",
            LLM_MODEL="deepseek-chat",
            LLM_BASE_URL="https://api.deepseek.com",
            LLM_HEALTH_CHECK_ENABLED=False,
        )

    app.dependency_overrides[get_app_settings] = override_settings
    client = TestClient(app)

    response = client.get("/health/llm")

    assert response.status_code == 200
    data = response.json()
    assert data["enabled"] is True
    assert data["ok"] is True
    assert "sk-secret" not in response.text
    assert data["model"] == "deepseek-chat"

    app.dependency_overrides.clear()


def test_audit_workbench_detail_table_layout_matches_review_workflow():
    """审批明细页面应该隐藏二级分类，并展示规则召回诊断列。"""
    client = TestClient(create_app())

    response = client.get("/")

    assert response.status_code == 200
    assert "<th>二级分类</th>" not in response.text
    assert "<th>召回粒度</th>" in response.text
    assert "<th>候选分类</th>" in response.text
    assert "<th>审批原因</th>" in response.text


def test_audit_workbench_uses_token_input_and_button_download():
    """前端应该通过 fetch 带认证头下载，而不是裸链接直接下载。"""
    client = TestClient(create_app())

    response = client.get("/")

    assert response.status_code == 200
    assert 'id="authToken"' in response.text
    assert 'id="downloadButton"' in response.text
    assert 'id="downloadLink"' not in response.text
    assert "/static/app.js?v=" in response.text


def test_frontend_download_fetches_excel_with_authorization_header():
    """前端下载 Excel 时，应该使用 Authorization header 和 blob 下载。"""
    app_js = (Path(__file__).resolve().parents[2] / "app" / "static" / "app.js").read_text(encoding="utf-8")

    assert "downloadResultExcel" in app_js
    assert "Authorization" in app_js
    assert "response.blob()" in app_js
    assert "URL.createObjectURL" in app_js
    assert "downloadLink.href = downloadUrl" not in app_js
    assert "getRequiredElement" in app_js


def test_frontend_validates_token_before_uploading():
    """前端未填写 token 时应该直接给出提示，不让用户误以为没反应。"""
    app_js = (Path(__file__).resolve().parents[2] / "app" / "static" / "app.js").read_text(encoding="utf-8")

    assert "请先填写访问令牌" in app_js
    assert "if (!getAuthToken())" in app_js


def test_upload_audit_excel_rejects_non_xlsx_file():
    """上传非 xlsx 文件时，应该返回 400。"""
    client = TestClient(create_app())

    response = client.post(
        "/api/audit/upload",
        files={
            "file": (
                "bad.txt",
                b"not excel",
                "text/plain",
            )
        },
    )

    assert response.status_code == 400
    assert "xlsx" in response.json()["detail"]


def test_upload_audit_excel_rejects_fake_xlsx_file(test_client):
    """伪造 .xlsx 后缀的文本文件应该被拒绝，而不是进入审批流程。"""
    response = test_client.post(
        "/api/audit/upload",
        files={
            "file": (
                "fake.xlsx",
                b"not a real xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 400
    assert "Excel 文件格式" in response.json()["detail"]


def _create_sample_excel(path):
    """创建测试用员工录入 Excel。"""
    workbook = Workbook()
    sheet = workbook.active

    sheet.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网"])
    sheet.append(
        [
            "2026-07-25",
            "张三",
            "信软技术服务",
            "软件开发",
            "杭州云启软件有限公司 https://yunqi.example.com",
        ]
    )

    workbook.save(path)


def test_download_audit_result_after_upload(tmp_path):
    """上传审批 Excel 后，应该可以下载结果 Excel。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    _create_sample_excel(excel_path)

    client = TestClient(create_app())

    with excel_path.open("rb") as file:
        upload_response = client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert upload_response.status_code == 200

    download_url = upload_response.json()["download_url"]
    download_response = client.get(download_url)

    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert download_response.content


def test_download_audit_result_records_access_audit_log(test_client, tmp_path):
    """下载结果 Excel 后，应该记录文件访问审计日志。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    _create_sample_excel(excel_path)

    with excel_path.open("rb") as file:
        upload_response = test_client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert upload_response.status_code == 200

    download_response = test_client.get(
        upload_response.json()["download_url"],
        headers={"User-Agent": "audit-test-client"},
    )

    assert download_response.status_code == 200
    logs = list_file_access_audit_log_payloads(tmp_path / "audit_agent.db")
    assert len(logs) == 1
    assert logs[0]["batch_id"] == upload_response.json()["batch_id"]
    assert logs[0]["action"] == "download_result_excel"

def test_download_audit_result_returns_404_when_missing():
    """下载不存在的批次结果时，应该返回 404。"""
    client = TestClient(create_app())

    response = client.get("/api/audit/batches/missing_batch/result")

    assert response.status_code == 404


def test_get_audit_batch_after_upload(tmp_path):
    """上传审批 Excel 后，应该能按 batch_id 查询批次摘要。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    _create_sample_excel(excel_path)

    client = TestClient(create_app())

    with excel_path.open("rb") as file:
        upload_response = client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert upload_response.status_code == 200

    batch_id = upload_response.json()["batch_id"]
    batch_response = client.get(f"/api/audit/batches/{batch_id}")

    assert batch_response.status_code == 200
    data = batch_response.json()
    assert data["batch_id"] == batch_id
    assert data["decision_count"] >= 1
    assert "result_file" not in data
    assert data["download_url"] == f"/api/audit/batches/{batch_id}/result"

    cases_response = client.get(f"/api/audit/batches/{batch_id}/cases")
    decisions_response = client.get(f"/api/audit/batches/{batch_id}/decisions")
    assessments_response = client.get(f"/api/audit/batches/{batch_id}/assessments")

    assert cases_response.status_code == 200
    assert decisions_response.status_code == 200
    assert assessments_response.status_code == 200
    assert cases_response.json()["total"] >= 1
    assert decisions_response.json()["total"] >= 1
    assert assessments_response.json()["total"] >= 1
    assert "classification_agent_status" in decisions_response.json()["items"][0]
    assert "classification_agent_error" in decisions_response.json()["items"][0]


def test_list_audit_batches_returns_items():
    """应该能查询审批批次列表。"""
    client = TestClient(create_app())

    response = client.get("/api/audit/batches")

    assert response.status_code == 200
    data = response.json()
    assert "total" in data
    assert "items" in data


def test_get_audit_batch_returns_404_when_missing():
    """查询不存在批次时，应该返回 404。"""
    client = TestClient(create_app())

    response = client.get("/api/audit/batches/not_exists")

    assert response.status_code == 404


def test_upload_uses_temp_storage(test_client, tmp_path):
    """上传接口应该使用测试注入的临时 storage。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    _create_sample_excel(excel_path)

    with excel_path.open("rb") as file:
        response = test_client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200

    data = response.json()
    assert "result_file" not in data
    assert (tmp_path / "audit_agent.db").exists()


def test_review_task_api_supports_query_and_manual_decision(test_client, tmp_path):
    """上传后应该能查询复核任务，并提交人工复核结论。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    _create_missing_evidence_excel(excel_path)

    with excel_path.open("rb") as file:
        upload_response = test_client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert upload_response.status_code == 200

    tasks_response = test_client.get("/api/review/tasks")
    assert tasks_response.status_code == 200
    tasks_data = tasks_response.json()
    assert tasks_data["total"] == 1

    task_id = tasks_data["items"][0]["task_id"]
    decision_response = test_client.post(
        f"/api/review/tasks/{task_id}/decision",
        json={
            "reviewer": "audit_user",
            "final_action": "approved",
        },
    )

    assert decision_response.status_code == 200
    assert decision_response.json()["status"] == "completed"
    assert decision_response.json()["reviewed_at"]


def test_download_result_excel_reflects_manual_review_decision(test_client, tmp_path):
    """人工复核提交后，再次下载 Excel 应该包含最新人工结论。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    _create_missing_evidence_excel(excel_path)

    with excel_path.open("rb") as file:
        upload_response = test_client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert upload_response.status_code == 200
    task = test_client.get("/api/review/tasks").json()["items"][0]

    decision_response = test_client.post(
        f"/api/review/tasks/{task['task_id']}/decision",
        json={
            "reviewer": "audit_user",
            "final_action": "approved",
        },
    )

    assert decision_response.status_code == 200

    download_response = test_client.get(upload_response.json()["download_url"])
    assert download_response.status_code == 200

    downloaded_path = tmp_path / "reviewed_result.xlsx"
    downloaded_path.write_bytes(download_response.content)
    workbook = load_workbook(downloaded_path)
    detail_sheet = workbook["审批明细"]
    headers = [cell.value for cell in detail_sheet[1]]
    final_action_column = headers.index("人工最终结论") + 1
    reviewer_column = headers.index("人工复核状态") + 1

    assert detail_sheet.cell(row=2, column=final_action_column).value == "approved"
    assert detail_sheet.cell(row=2, column=reviewer_column).value == "completed"


def test_upload_audit_excel_keeps_invalid_rows_as_review_tasks(test_client, tmp_path):
    """上传真实 Excel 时，单行缺字段不应该导致整个批次失败。"""
    excel_path = tmp_path / "audit_cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网"])
    sheet.append(["2026-07-25", "张三", "信软技术服务", "软件开发", "杭州云启软件有限公司"])
    sheet.append(["2026-07-25", "", "信软技术服务", "软件开发", "杭州云启软件有限公司"])
    workbook.save(excel_path)

    with excel_path.open("rb") as file:
        response = test_client.post(
            "/api/audit/upload",
            files={
                "file": (
                    "audit_cases.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200
    data = response.json()
    assert data["case_count"] == 2
    assert data["decision_count"] == 2
    assert any("缺少必填字段" in decision["reason"] for decision in data["decisions"])


def _create_missing_evidence_excel(path):
    """创建一条证据库中不存在企业的 Excel，用于触发人工复核。"""
    workbook = Workbook()
    sheet = workbook.active

    sheet.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网"])
    sheet.append(
        [
            "2026-07-25",
            "张三",
            "信软技术服务",
            "软件开发",
            "不存在的样例企业有限公司 https://missing.example.com",
        ]
    )

    workbook.save(path)
