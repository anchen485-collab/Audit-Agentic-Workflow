from pathlib import Path

from openpyxl import Workbook

from app.domain.models import AgentAssessment, ApprovalDecision, AuditCase, EvidencePackage, ReviewTask


def write_audit_result_excel(
    output_path: str | Path,
    cases: list[AuditCase],
    decisions: list[ApprovalDecision],
    review_tasks: list[ReviewTask],
    evidence_packages: dict[str, EvidencePackage] | None = None,
    assessments: dict[str, list[AgentAssessment]] | None = None,
) -> Path:
    """导出审批结果 Excel。"""
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    detail_sheet = workbook.active
    detail_sheet.title = "审批明细"

    _write_detail_sheet(detail_sheet, cases, decisions, review_tasks, evidence_packages or {}, assessments or {})
    _write_summary_sheet(workbook, cases, decisions)
    _write_review_sheet(workbook, review_tasks)

    workbook.save(output_file)
    return output_file


def _write_detail_sheet(
    sheet,
    cases: list[AuditCase],
    decisions: list[ApprovalDecision],
    review_tasks: list[ReviewTask],
    evidence_packages: dict[str, EvidencePackage],
    assessments: dict[str, list[AgentAssessment]],
) -> None:
    """写入审批明细 Sheet。"""
    decision_map = {decision.case_id: decision for decision in decisions}
    review_task_map = {task.case_id: task for task in review_tasks}
    sheet.append(
        [
            "审批单ID",
            "批次ID",
            "原始行号",
            "录入人员",
            "企业名称",
            "企业官网",
            "员工录入一级分类",
            "员工录入二级分类",
            "员工录入细分",
            "细分粒度",
            "候选分类",
            "召回原因",
            "审批动作",
            "审批状态",
            "置信度",
            "风险分",
            "审批原因",
            "处理建议",
            "是否需要人工复核",
            "人工复核状态",
            "人工最终结论",
            "人工复核时间",
            "证据来源",
            "证据质量分",
            "证据片段",
            "证据错误",
            "Classification Agent 状态",
            "Classification Agent 错误原因",
            "证据不足类型",
        ]
    )

    for case in cases:
        decision = decision_map.get(case.case_id)
        review_task = review_task_map.get(case.case_id)
        evidence = evidence_packages.get(case.case_id)
        classification_assessment = _find_agent_assessment(
            assessments.get(case.case_id, []),
            "classification_agent",
        )
        rule_assessment = _find_agent_assessment(
            assessments.get(case.case_id, []),
            "rule_agent",
        )
        sheet.append(
            [
                case.case_id,
                case.batch_id,
                case.row_number,
                case.employee_name,
                case.company_name,
                case.website_url,
                case.entered_level1,
                case.entered_level2,
                case.entered_subcategory,
                _rule_granularity(rule_assessment),
                _rule_candidate_paths(rule_assessment),
                _rule_granularity_reason(rule_assessment),
                decision.action.value if decision else "",
                decision.status.value if decision else "",
                decision.confidence if decision else "",
                decision.risk_score if decision else "",
                decision.reason if decision else "",
                decision.suggestion if decision else "",
                decision.needs_review if decision else "",
                review_task.status if review_task else "",
                review_task.final_action if review_task else "",
                review_task.reviewed_at if review_task else "",
                "、".join(evidence.sources) if evidence else "",
                evidence.quality_score if evidence else "",
                " | ".join(snippet.text for snippet in evidence.snippets) if evidence else "",
                " | ".join(evidence.errors) if evidence else "",
                classification_assessment.status if classification_assessment else "",
                _agent_error_reason(classification_assessment),
                _evidence_error_type(classification_assessment),
            ]
        )


def _write_summary_sheet(workbook: Workbook, cases: list[AuditCase], decisions: list[ApprovalDecision]) -> None:
    """写入人员汇总 Sheet。"""
    sheet = workbook.create_sheet("人员汇总")
    decision_map = {decision.case_id: decision for decision in decisions}
    summary: dict[str, dict[str, int]] = {}

    for case in cases:
        person = summary.setdefault(
            case.employee_name,
            {
                "total": 0,
                "approved": 0,
                "rejected": 0,
                "needs_more_evidence": 0,
                "needs_human_review": 0,
            },
        )
        person["total"] += 1
        decision = decision_map.get(case.case_id)
        if decision:
            person[decision.action.value] += 1

    sheet.append(["录入人员", "总数", "自动通过", "自动驳回", "证据不足", "人工复核"])
    for employee_name, values in summary.items():
        sheet.append(
            [
                employee_name,
                values["total"],
                values["approved"],
                values["rejected"],
                values["needs_more_evidence"],
                values["needs_human_review"],
            ]
        )


def _write_review_sheet(workbook: Workbook, review_tasks: list[ReviewTask]) -> None:
    """写入人工复核任务 Sheet。"""
    sheet = workbook.create_sheet("人工复核任务")
    sheet.append(["复核任务ID", "审批单ID", "优先级", "复核原因", "任务状态", "审核员", "人工最终动作", "人工复核时间"])

    for task in review_tasks:
        sheet.append(
            [
                task.task_id,
                task.case_id,
                task.priority,
                task.reason,
                task.status,
                task.reviewer or "",
                task.final_action or "",
                task.reviewed_at or "",
            ]
        )


def _find_agent_assessment(
    assessments: list[AgentAssessment],
    agent_name: str,
) -> AgentAssessment | None:
    """从某条 case 的 assessments 中找到指定 Agent 的执行记录。"""
    for assessment in assessments:
        if assessment.agent_name == agent_name:
            return assessment

    return None


def _agent_error_reason(assessment: AgentAssessment | None) -> str:
    """提取 Agent 错误或降级原因。"""
    if assessment is None:
        return ""

    if assessment.status == "failed":
        return str(assessment.output.get("error", ""))

    return str(assessment.output.get("normalization_warning", ""))


def _evidence_error_type(assessment: AgentAssessment | None) -> str:
    """提取证据不足类型，方便人工快速区分官网失败和内容不足。"""
    if assessment is None:
        return ""

    return str(assessment.output.get("evidence_error_type", ""))


def _rule_granularity(assessment: AgentAssessment | None) -> str:
    """提取 Rule Agent 识别出的员工细分粒度。"""
    if assessment is None:
        return ""

    return str(assessment.output.get("entered_subcategory_granularity", ""))


def _rule_candidate_paths(assessment: AgentAssessment | None) -> str:
    """提取 Rule Agent 召回的分类候选路径。"""
    if assessment is None:
        return ""

    paths = assessment.output.get("category_candidate_paths", [])
    if not isinstance(paths, list):
        return ""

    return " | ".join(str(path) for path in paths)


def _rule_granularity_reason(assessment: AgentAssessment | None) -> str:
    """提取 Rule Agent 的粒度识别原因。"""
    if assessment is None:
        return ""

    return str(assessment.output.get("granularity_reason", ""))
