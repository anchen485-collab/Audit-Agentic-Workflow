import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, Field

REQUIRED_HEADERS = ["日期", "姓名", "一级分类", "细分", "企业名称&官网"]
OPTIONAL_HEADER_ALIASES = {
    "二级分类": ["二级分类", "二级品类", "二级"],
}
URL_PATTERN = re.compile(r"(https?://\S+|www\.\S+)", re.IGNORECASE)


class MissingExcelHeaderError(ValueError):
    """员工录入 Excel 缺少必要表头。"""


class ExcelAuditRow(BaseModel):
    """员工录入 Excel 中的一行原始数据。"""

    row_number: int = Field(description="Excel 原始行号")
    date: str = Field(default="", description="录入日期")
    employee_name: str = Field(default="", description="录入员工姓名")
    entered_level1: str = Field(default="", description="员工录入一级分类")
    entered_level2: str = Field(default="", description="员工录入二级分类")
    entered_subcategory: str = Field(default="", description="员工录入细分")
    company_raw: str = Field(default="", description="Excel 中原始企业字段")
    company_name: str = Field(default="", description="清洗后的企业名称")
    website_url: str = Field(default="", description="企业官网 URL")


def read_audit_excel(file_path: str | Path) -> list[ExcelAuditRow]:
    """读取员工录入 Excel，并返回结构化原始行。"""
    workbook = load_workbook(filename=file_path, data_only=True)
    sheet = workbook.active

    header_cells = next(sheet.iter_rows(min_row=1, max_row=1))
    header_values = [_cell_to_text(cell.value) for cell in header_cells]
    header_map = _build_header_map(header_values)

    rows: list[ExcelAuditRow] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = tuple(cell.value for cell in cells)
        if _is_empty_row(values):
            continue

        company_cell = cells[header_map["企业名称&官网"]]
        company_raw = _cell_to_text(company_cell.value)
        company_name, parsed_url = parse_company_field(company_raw)
        website_url = parsed_url or _cell_hyperlink_target(company_cell)

        rows.append(
            ExcelAuditRow(
                row_number=row_number,
                date=_cell_to_text(values[header_map["日期"]]),
                employee_name=_cell_to_text(values[header_map["姓名"]]),
                entered_level1=_cell_to_text(values[header_map["一级分类"]]),
                entered_level2=_optional_cell_text(values, header_map, "二级分类"),
                entered_subcategory=_cell_to_text(values[header_map["细分"]]),
                company_raw=company_raw,
                company_name=company_name,
                website_url=website_url,
            )
        )

    return rows


def parse_company_field(value: Any) -> tuple[str, str]:
    """从企业字段中拆出企业名称和官网 URL。"""
    text = _cell_to_text(value)
    match = URL_PATTERN.search(text)
    if not match:
        return _compact_text(text), ""

    website_url = _normalize_url(match.group(1))
    company_name = _compact_text(URL_PATTERN.sub("", text).strip("，,；;。 "))

    return company_name, website_url


def _cell_hyperlink_target(cell) -> str:
    """读取单元格隐藏超链接，兼容企业名显示文本和官网链接分离的 Excel。"""
    if cell.hyperlink is None:
        return ""

    return _normalize_url(cell.hyperlink.target)


def _build_header_map(headers: list[str]) -> dict[str, int]:
    """构建表头到列索引的映射，并校验必填表头。"""
    header_map = {header: index for index, header in enumerate(headers) if header}
    missing_headers = [header for header in REQUIRED_HEADERS if header not in header_map]
    if missing_headers:
        raise MissingExcelHeaderError(f"员工录入 Excel 缺少必要表头：{', '.join(missing_headers)}")
    for standard_header, aliases in OPTIONAL_HEADER_ALIASES.items():
        for alias in aliases:
            if alias in header_map:
                header_map[standard_header] = header_map[alias]
                break
    return header_map


def _optional_cell_text(values: tuple[Any, ...], header_map: dict[str, int], header_name: str) -> str:
    """读取可选列；Excel 没有这列时返回空字符串。"""
    index = header_map.get(header_name)
    if index is None or index >= len(values):
        return ""
    return _cell_to_text(values[index])


def _is_empty_row(values: tuple[Any, ...]) -> bool:
    """判断一行是否完全为空。"""
    return all(_cell_to_text(value) == "" for value in values)


def _cell_to_text(value: Any) -> str:
    """把 Excel 单元格值统一转成字符串。"""
    if value is None:
        return ""
    return str(value).strip()


def _normalize_url(value: str) -> str:
    """标准化 URL，支持 www 开头的网址。"""
    raw_url = _cell_to_text(value).strip("，,；;。 ")
    if not raw_url:
        return ""

    if raw_url.startswith(("http://", "https://")):
        return raw_url

    return f"https://{raw_url}"


def _compact_text(value: str) -> str:
    """压缩多余空白，避免企业名中残留换行或多个空格。"""
    return re.sub(r"\s+", " ", value or "").strip()
