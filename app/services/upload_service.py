from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import is_zipfile

from fastapi import UploadFile
from openpyxl import load_workbook

from app.infrastructure.settings import Settings

ALLOWED_EXCEL_SUFFIXES = {".xlsx"}
MAX_UPLOAD_SIZE_BYTES = 5 * 1024 * 1024


async def save_upload_file(file: UploadFile, settings: Settings) -> tuple[str, Path]:
    """保存上传的 Excel 文件，并返回 batch_id 和保存路径。"""
    original_filename = file.filename or ""

    if not original_filename:
        raise ValueError("上传文件名不能为空。")

    suffix = Path(original_filename).suffix.lower()
    if suffix not in ALLOWED_EXCEL_SUFFIXES:
        raise ValueError("只支持上传 .xlsx 文件。")

    batch_id = f"batch_{uuid4().hex[:12]}"
    upload_path = settings.uploads_dir / f"{batch_id}{suffix}"
    upload_path.parent.mkdir(parents=True, exist_ok=True)

    content = await file.read()
    if not content:
        raise ValueError("上传文件内容不能为空。")

    if len(content) > MAX_UPLOAD_SIZE_BYTES:
        raise ValueError("上传文件不能超过 5MB。")

    _validate_xlsx_content(content=content, settings=settings)

    upload_path.write_bytes(content)

    return batch_id, upload_path


def _validate_xlsx_content(content: bytes, settings: Settings) -> None:
    """校验上传内容确实是 xlsx，并限制工作簿规模，避免异常大文件拖垮服务。"""
    content_stream = BytesIO(content)
    if not is_zipfile(content_stream):
        raise ValueError("Excel 文件格式不正确，请上传真实的 .xlsx 文件。")

    content_stream.seek(0)
    try:
        workbook = load_workbook(content_stream, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("Excel 文件格式不正确，请上传可读取的 .xlsx 文件。") from error

    try:
        if len(workbook.worksheets) > settings.audit_max_excel_sheets:
            raise ValueError(f"Excel sheet 数不能超过 {settings.audit_max_excel_sheets} 个。")

        for sheet in workbook.worksheets:
            if sheet.max_row > settings.audit_max_excel_rows:
                raise ValueError(f"Excel 行数不能超过 {settings.audit_max_excel_rows} 行。")
            if sheet.max_column > settings.audit_max_excel_columns:
                raise ValueError(f"Excel 列数不能超过 {settings.audit_max_excel_columns} 列。")
    finally:
        workbook.close()
